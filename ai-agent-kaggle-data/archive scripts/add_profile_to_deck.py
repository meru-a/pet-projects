from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).parent
SOURCE = ROOT / "Customer_Marketing_Executive_Analysis_revised.pptx"
OUTPUT = ROOT / "Customer_Marketing_Executive_Analysis_revised_final.pptx"

NAVY = RGBColor(15, 39, 71)
TEAL = RGBColor(26, 150, 136)
INK = RGBColor(40, 52, 66)
MUTED = RGBColor(105, 119, 135)

prs = Presentation(SOURCE)
slide = prs.slides[1]

def set_text(shape, value, size, color, bold=False):
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]; p.text = value
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color

# Reuse the lower-slide panel as a compact persona definition for the high-value segment.
set_text(slide.shapes[20], "HIGH-VALUE CUSTOMER PROFILE", 10, TEAL, True)
slide.shapes[20].left = Inches(.82); slide.shapes[20].top = Inches(5.76); slide.shapes[20].width = Inches(3.0); slide.shapes[20].height = Inches(.18)

set_text(slide.shapes[21], "TOP SPEND QUARTILE  |  555 customers  |  €1,048+ total spend", 11, NAVY, True)
slide.shapes[21].left = Inches(.82); slide.shapes[21].top = Inches(6.01); slide.shapes[21].width = Inches(5.35); slide.shapes[21].height = Inches(.18)

def add_text(value, x, y, w, h, size, color, bold=False):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(sh, value, size, color, bold)
    return sh

add_text("WHO THEY ARE", 6.70, 5.76, 1.35, .18, 10, TEAL, True)
add_text("€75k average income  |  93% degree+  |  33% with children/teens at home", 6.70, 6.01, 5.35, .18, 10, INK, True)
add_text("WHAT THEY BUY", .82, 6.34, 1.35, .18, 10, TEAL, True)
add_text("€741 wine + €455 meat = 80% of basket; meat spend is 6.4× the rest of the base.", 2.20, 6.34, 4.30, .22, 9.5, MUTED)
add_text("HOW TO ACT", 6.70, 6.34, 1.05, .18, 10, TEAL, True)
add_text("Lead with premium wine-and-meat bundles; prioritize prior campaign accepters.", 7.88, 6.34, 4.12, .22, 9.5, MUTED)

# Update the notes so the talk track introduces the newly visible target persona.
slide.notes_slide.notes_text_frame.text = (
    "Response is not uniform across the file: high-spend customers are more responsive, and prior acceptance creates a much stronger lift. "
    "The highest-spend quartile is a clear persona: higher-income, highly educated, low-child-household customers whose basket is dominated by wine and meat. "
    "Use that profile to shape creative and contact investment, while retaining a holdout to validate incremental profit."
)

prs.save(OUTPUT)
print(f"Created {OUTPUT}")
