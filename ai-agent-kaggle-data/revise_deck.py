from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

ROOT = Path(__file__).parent
SOURCE = ROOT / "Customer_Marketing_Executive_Analysis.pptx"
OUTPUT = ROOT / "Customer_Marketing_Executive_Analysis_revised.pptx"

NAVY = RGBColor(15, 39, 71)
TEAL = RGBColor(26, 150, 136)
BLUE = RGBColor(32, 104, 170)
GOLD = RGBColor(233, 162, 45)
INK = RGBColor(40, 52, 66)
MUTED = RGBColor(105, 119, 135)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(228, 240, 249)

prs = Presentation(SOURCE)
slide = prs.slides[0]

def rect(x, y, w, h, fill, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    return shape

def label(value, x, y, w, h, size, color, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]; p.text = value
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    if align is not None: p.alignment = align
    return shape

# Preserve the user's title/subtitle edit; only rebalance the content below the divider.
# Existing KPI cards, chart, chart interpretation, and insight panel move down as a coherent grid.
for i in [5, 6, 7]: slide.shapes[i].top = Inches(2.60 if i == 5 else (2.73 if i == 6 else 3.11))
for i in [8, 9, 10]: slide.shapes[i].top = Inches(2.60 if i == 8 else (2.73 if i == 9 else 3.11))
for i in [11, 12, 13]: slide.shapes[i].top = Inches(2.60 if i == 11 else (2.73 if i == 12 else 3.11))
slide.shapes[14].top = Inches(2.54)
slide.shapes[15].top = Inches(4.88)
slide.shapes[16].top = Inches(3.86); slide.shapes[16].height = Inches(2.86)
slide.shapes[17].top = Inches(4.10)
for idx, top in zip(range(18, 23), [4.52, 4.91, 5.30, 5.69, 6.08]):
    slide.shapes[idx].top = Inches(top)

# Dataset context comes immediately after the title divider and before any findings.
rect(.55, 1.68, 12.18, .62, WHITE, True)
rect(.75, 1.84, .06, .27, TEAL, True)
label("DATASET AT A GLANCE", .95, 1.79, 1.85, .17, 9, TEAL, True)
items = [
    ("2,216", "customers with usable income data"),
    ("2 years", "of spend and purchase history"),
    ("6 categories", "product-level spend signals"),
    ("3 channels", "web, catalogue and store purchases"),
]
for j, (big, small) in enumerate(items):
    x = 3.05 + j * 2.35
    label(big, x, 1.78, 1.05, .20, 13, NAVY, True)
    label(small, x, 2.02, 2.12, .16, 7.5, MUTED)
    if j < 3:
        rect(x + 2.10, 1.82, .012, .29, RGBColor(218, 227, 235))

# Use the former empty lower-right area for the slide-level executive takeaway.
rect(6.95, 5.37, 5.72, 1.20, NAVY, True)
label("EXECUTIVE READ-THROUGH", 7.22, 5.63, 2.20, .16, 9, RGBColor(173, 214, 238), True)
label("Focus investment where value and response overlap: protect premium spenders, then use campaign memory to improve contact efficiency.", 7.22, 5.91, 4.92, .43, 11, WHITE, True)

prs.save(OUTPUT)
print(f"Created {OUTPUT}")
